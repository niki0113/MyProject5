"""視覚メディア研究室の発表スライドを組み立てる部品集（仁木デザイン）。

研究室の発表資料から「わかりやすさの仕組み」（章名の常時表示、矢印から結論への流れ、
番号つきの枠、注目列を色分けした表）を借り、見た目は独自に組んだもの。
配色は仁木さんの 7/28 の資料（青と黄色の下線）に合わせている。

使い方:
    import sys; sys.path.insert(0, "<スキルの scripts ディレクトリ>")
    from vml_deck import Deck, Box
    d = Deck(chapters=["研究背景", "提案手法", "進捗", "今後"])
    d.title_slide("大ゼミ（2026/10/28）", ["一行目", "二行目"], "S0000000　氏名")
    s = d.content("見出し", chapter=0)
    d.bullets(s, ["項目 1", "項目 2"])
    d.save("out.pptx")

座標はすべて EMU。キャンバスは 12192000 x 6858000（13.33 x 7.5 インチ）。
本文領域は Box.BODY（見出しの下の区切り線から、下端のページ表示の上まで）。
"""
import copy
import os
from dataclasses import dataclass

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
TEMPLATE = os.path.join(ASSETS, "template.pptx")
LOGO = os.path.join(ASSETS, "vmlab_logo.png")

FONT = "BIZ UDPゴシック"   # Windows 10 以降の標準。UD フォントで、和欧とも読みやすい

# 色（仁木さんの 7/28 の資料の青を基準にした濃淡と、決め文の黄色）
BLUE = RGBColor(0x4E, 0x95, 0xD9)        # メイン：章の札、行頭の点、キーワード、下端の線
NAVY = RGBColor(0x1F, 0x3A, 0x5F)        # 見出しの文字、表の見出し行、番号バッジ
BLUE_DARK = RGBColor(0x21, 0x5F, 0x9A)   # 強調語の文字、ラベルの文字
PALE = RGBColor(0xDC, 0xEA, 0xF7)        # 結論の面、ラベルの札、注目する列
MIST = RGBColor(0xF3, 0xF8, 0xFD)        # 流れ図の箱、選択中の枠の地
YELLOW = "FFC000"                          # 最重要語の太い下線
TEXT = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)        # 補足、下端の文字、矢印
HAIR = RGBColor(0xD9, 0xDE, 0xE6)        # 区切り線、枠の細線
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 文字サイズ（pt）。標準の 16:9 キャンバス
SZ_TITLE = 30
SZ_KEY = 24
SZ_BODY = 22
SZ_SUB = 20
SZ_SMALL = 18     # 本文の下限（カード内、表、図のラベル）
SZ_FOOT = 11      # 下端の表示と参考文献だけ

W, H = 12192000, 6858000
IN = 914400


@dataclass
class Box:
    x: int
    y: int
    w: int
    h: int

    BODY = None

    @property
    def r(self):
        return self.x + self.w

    @property
    def b(self):
        return self.y + self.h


Box.BODY = Box(int(0.6 * IN), int(1.5 * IN), int(12.13 * IN), int(5.0 * IN))


def _font(r, size=None, bold=None, color=None):
    f = r.font
    f.name = FONT
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


def _flat(shp):
    shp.shadow.inherit = False
    return shp


def _rect(slide, kind, box, fill=None, line=None, line_w=1.0, radius=None, name=None):
    shp = _flat(slide.shapes.add_shape(kind, box.x, box.y, box.w, box.h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    if radius is not None:
        shp.adjustments[0] = radius
    if name:
        shp.name = name
    return shp


class Deck:
    def __init__(self, chapters=(), template=TEMPLATE):
        self.prs = Presentation(template)
        # 元の資料の変更履歴（編集者の氏名とアカウント）、改訂情報、SharePoint の管理情報を外す
        for rId, rel in list(self.prs.part.rels.items()):
            if any(k in rel.reltype for k in ("changesInfo", "revisionInfo", "/customXml")):
                self.prs.part.drop_rel(rId)
        self.prs.slide_width, self.prs.slide_height = W, H
        self._blank = self.prs.slide_layouts[6]  # 白紙
        self.chapters = list(chapters)
        self._pages = []  # (スライド, 番号の図形)。保存時に「n / N」を埋める

    # ------------------------------------------------------------ スライドの枠
    def _slide(self, notes=None, chrome=True):
        s = self.prs.slides.add_slide(self._blank)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        if chrome:
            self._chrome(s)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def _chrome(self, s):
        """下端：最下部に細い青の線、右下に「n / N」（保存時に埋める）。"""
        num = self.text(s, Box(int(10.73 * IN), int(7.0 * IN), int(2.0 * IN), int(0.25 * IN)), [""],
                        size=SZ_FOOT + 1, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        _rect(s, MSO_SHAPE.RECTANGLE, Box(0, H - int(0.06 * IN), W, int(0.06 * IN)), fill=BLUE, name="下端の線")
        self._pages.append((s, num))

    def _header(self, s, title, chapter=None):
        """左上に章の札（番号と章名）、その下に左揃えの見出し、下に細い区切り線。"""
        if chapter is not None:
            label = chapter if isinstance(chapter, str) else f"{chapter + 1:02d}  {self.chapters[chapter]}"
            w = int(Pt(13) * (len(label) * 0.9 + 2.4))
            pill = _rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, Box(int(0.6 * IN), int(0.32 * IN), w, int(0.3 * IN)),
                         fill=BLUE, radius=0.5, name="章の札")
            tf = pill.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = label; _font(r, 13, True, WHITE)
        runs = [(title, {})] if isinstance(title, str) else title
        self.text(s, Box(int(0.6 * IN), int(0.68 * IN), int(12.1 * IN), int(0.6 * IN)), [runs], size=SZ_TITLE,
                  bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
        _rect(s, MSO_SHAPE.RECTANGLE, Box(int(0.6 * IN), int(1.34 * IN), int(12.13 * IN), 12700), fill=HAIR,
              name="区切り線")

    def title_slide(self, event, title_lines, name, notes=None, image=None, image_label="研究テーマの画像"):
        """表紙。右上に月を思わせる大きな淡い円、左に行事名の札、タイトル、氏名、ロゴ。右に研究テーマの画像。"""
        s = self._slide(notes)
        moon = _rect(s, MSO_SHAPE.OVAL, Box(int(7.6 * IN), int(-1.6 * IN), int(7.2 * IN), int(7.2 * IN)), fill=MIST,
                     name="背景の円")
        sp = s.shapes._spTree; el = moon._element; sp.remove(el); sp.insert(2, el)
        w = int(Pt(14) * (len(event) * 0.9 + 2.6))
        pill = _rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, Box(int(0.7 * IN), int(1.0 * IN), w, int(0.36 * IN)),
                     line=BLUE, line_w=1.25, fill=WHITE, radius=0.5, name="行事名の札")
        tf = pill.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = event; _font(r, 14, True, BLUE_DARK)
        self.text(s, Box(int(0.7 * IN), int(1.6 * IN), int(7.3 * IN), int(1.9 * IN)), title_lines, size=32, bold=True,
                  color=NAVY, line_spacing=1.15, anchor=MSO_ANCHOR.TOP)
        self.text(s, Box(int(0.7 * IN), int(3.85 * IN), int(6 * IN), int(0.4 * IN)), [name], size=SZ_SUB, color=TEXT)
        s.shapes.add_picture(LOGO, int(0.7 * IN), int(5.55 * IN), height=int(0.85 * IN))
        self.image(s, Box(int(8.25 * IN), int(2.35 * IN), int(4.4 * IN), int(3.3 * IN)), path=image, label=image_label)
        return s

    def section(self, chapter, notes=None):
        """章扉。左に大きな薄青の章番号と章名、右に目次（今の章だけ強調）。"""
        s = self._slide(notes)
        self.text(s, Box(int(0.5 * IN), int(1.2 * IN), int(5.5 * IN), int(2.4 * IN)),
                  [[(f"{chapter + 1:02d}", {"bold": True, "color": PALE})]], size=150, line_spacing=0.9,
                  anchor=MSO_ANCHOR.BOTTOM)
        self.text(s, Box(int(0.75 * IN), int(3.75 * IN), int(6.2 * IN), int(1.4 * IN)),
                  [[(self.chapters[chapter], {"bold": True, "color": NAVY})]], size=36, anchor=MSO_ANCHOR.TOP)
        _rect(s, MSO_SHAPE.RECTANGLE, Box(int(7.35 * IN), int(1.6 * IN), 12700, int(4.0 * IN)), fill=HAIR)
        self.agenda(s, self.chapters, Box(int(7.75 * IN), int(1.9 * IN), int(5.0 * IN), int(3.6 * IN)),
                    current=chapter, size=SZ_SUB)
        return s

    def content(self, title, chapter=None, notes=None):
        s = self._slide(notes)
        self._header(s, title, chapter)
        return s

    # ------------------------------------------------------------ 部品
    def agenda(self, slide, items, box=None, current=None, size=SZ_BODY + 2):
        """目次。青い番号と項目名。current（0 始まり）を渡すと、その項目だけを濃く、ほかを灰色にする。"""
        box = box or Box(int(3.2 * IN), int(2.0 * IN), int(7 * IN), int(4.0 * IN))
        lines = []
        for i, t in enumerate(items):
            on = current is None or i == current
            lines.append([(f"{i + 1:02d}", {"bold": True, "color": BLUE if on else HAIR}),
                          ("　" + t, {"bold": on, "color": NAVY if on else GRAY})])
        return self.text(slide, box, lines, size=size, line_spacing=1.0, space_after=size * 0.9)

    def bullets(self, slide, items, box=None, size=SZ_BODY, gap=None, sub_size=SZ_SMALL):
        """行頭に青い丸を付けた項目。items は [行, 行, ...]。1 項目を複数行にするときは行のリストにする。
        {"sub": [...]} を項目に含めると、その下に灰色寄りの小さい 2 段目を付ける。
        各行は文字列か run のリスト。返り値は最後の項目の下端 y。"""
        box = box or Box(Box.BODY.x, Box.BODY.y + 100000, Box.BODY.w, Box.BODY.h)
        gap = gap if gap is not None else int(Pt(size) * 0.65)
        y = box.y
        dot = int(Pt(size) * 0.36)
        for item in items:
            lines = item if isinstance(item, list) and item and not isinstance(item[0], tuple) else [item]
            main, subs = [], []
            for ln in lines:
                if isinstance(ln, dict):
                    subs.extend(ln["sub"])
                else:
                    main.append(ln)
            lh = int(Pt(size) * 1.45)
            _rect(slide, MSO_SHAPE.OVAL, Box(box.x, y + lh // 2 - dot // 2, dot, dot), fill=BLUE, name="行頭の点")
            tx = box.x + int(Pt(size) * 0.85)
            self.text(slide, Box(tx, y, box.r - tx, lh * len(main)), main, size=size, color=TEXT, line_spacing=1.2,
                      anchor=MSO_ANCHOR.MIDDLE if len(main) == 1 else MSO_ANCHOR.TOP)
            y += lh * len(main)
            if subs:
                sh = int(Pt(sub_size) * 1.45) * len(subs)
                self.text(slide, Box(tx, y, box.r - tx, sh), subs, size=sub_size, color=GRAY, line_spacing=1.2)
                y += sh
            y += gap
        return y

    def key_message(self, slide, runs, y, size=SZ_KEY, x=None, w=None, h=None):
        """結論の一文。淡い青の角丸の面に太字で置く。最重要語は {"color": "key"} で黄色の太い下線。"""
        x = Box.BODY.x if x is None else x
        w = Box.BODY.w if w is None else w
        lines = runs if runs and isinstance(runs[0], list) else [runs]
        h = h or int(Pt(size) * 1.5 * len(lines)) + int(0.32 * IN)
        _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Box(x, y, w, h), fill=PALE, radius=0.18, name="結論の面")
        return self.text(slide, Box(x + int(0.3 * IN), y, w - int(0.6 * IN), h), lines, size=size, bold=True,
                         color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def arrow_down(self, slide, cx, y, w=None, h=None):
        """帰結を示す下向きの山形（灰色）。"""
        w = w or int(0.6 * IN); h = h or int(0.32 * IN)
        a = _rect(slide, MSO_SHAPE.CHEVRON, Box(cx - h // 2, y - (w - h) // 2, h, w), fill=GRAY, name="帰結の矢印")
        a.rotation = 90
        return a

    def pill(self, slide, box, text, size=SZ_SMALL, fill=PALE, color=None):
        """角丸の札（ラベル）。"""
        b = _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, box, fill=fill, radius=0.5)
        tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
        tf.margin_left = tf.margin_right = 45720; tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text; _font(r, size, True, color or BLUE_DARK)
        return b

    def badge(self, slide, cx, cy, text, d=None, fill=None, size=16):
        """丸い番号バッジ。"""
        d = d or int(0.48 * IN)
        c = _rect(slide, MSO_SHAPE.OVAL, Box(cx - d // 2, cy - d // 2, d, d), fill=fill or NAVY)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text; _font(r, size, True, WHITE)
        return c

    def num_cards(self, slide, cards, box, gap=None, selected=None):
        """角丸の枠を横に並べる。上に丸い番号バッジ、名前、説明、下に札。
        cards は [(番号, 名前, 説明の行 or None, 札の語 or None), ...]。selected の枠は青い枠線と淡い地。"""
        n = len(cards)
        gap = gap if gap is not None else int(0.25 * IN)
        cw = (box.w - gap * (n - 1)) // n
        for i, (num, name, desc, tag) in enumerate(cards):
            x = box.x + i * (cw + gap)
            sel = selected == i
            _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Box(x, box.y, cw, box.h), fill=MIST if sel else WHITE,
                  line=BLUE if sel else HAIR, line_w=2.25 if sel else 1.25, radius=0.08)
            y = box.y + int(0.2 * IN)
            if num:
                if len(num) <= 2:
                    self.badge(slide, x + cw // 2, y + int(0.24 * IN), num, fill=BLUE if sel else NAVY)
                    y += int(0.6 * IN)
                else:
                    self.text(slide, Box(x, y, cw, int(0.35 * IN)), [[(num, {"bold": True, "color": BLUE})]],
                              size=SZ_SMALL, align=PP_ALIGN.CENTER)
                    y += int(0.4 * IN)
            self.text(slide, Box(x, y, cw, int(0.45 * IN)), [[(name, {"bold": True})]], size=SZ_BODY + 2,
                      align=PP_ALIGN.CENTER, color=NAVY)
            y += int(0.55 * IN)
            if desc:
                dl = desc if isinstance(desc, list) else [desc]
                dh = int(Pt(SZ_SMALL) * 1.45) * len(dl)
                self.text(slide, Box(x + 60000, y, cw - 120000, dh), dl, size=SZ_SMALL, align=PP_ALIGN.CENTER,
                          color=TEXT)
            if tag:
                self.pill(slide, Box(x + int(0.2 * IN), box.b - int(0.62 * IN), cw - int(0.4 * IN), int(0.42 * IN)),
                          tag, fill=BLUE if sel else PALE, color=WHITE if sel else None)

    def flow(self, slide, items, box, size=SZ_SMALL, gap=None, selected=None, numbered=True):
        """横に並ぶ処理の順序。淡い地の角丸の箱に番号バッジを付け、灰色の山形でつなぐ。
        selected の箱は青い地に白文字（今いる段階）。"""
        n = len(items)
        gap = gap or int(0.42 * IN)
        bw = (box.w - gap * (n - 1)) // n
        pos = []
        for i, t in enumerate(items):
            b = Box(box.x + i * (bw + gap), box.y, bw, box.h)
            pos.append(b)
            sel = selected == i
            shp = _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, b, fill=BLUE if sel else MIST,
                        line=None if sel else HAIR, radius=0.14)
            tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_top = int(0.18 * IN) if numbered else 45720
            lines = t if isinstance(t, list) else [t]
            for j, line in enumerate(lines):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                self._runs(p, line, size=size, bold=(j == 0), color=WHITE if sel else (NAVY if j == 0 else TEXT))
            if numbered:
                self.badge(slide, b.x + b.w // 2, b.y, str(i + 1), d=int(0.36 * IN), size=13,
                           fill=BLUE_DARK if sel else BLUE)
            if i < n - 1:
                aw = int(0.2 * IN)
                _rect(slide, MSO_SHAPE.CHEVRON, Box(b.r + (gap - aw) // 2, b.y + b.h // 2 - int(0.16 * IN), aw,
                                                    int(0.32 * IN)), fill=GRAY)
        return pos

    def table(self, slide, box, rows, col_w=None, size=SZ_SMALL, highlight_col=None, row_h=None,
              align=None, bold_rows=()):
        """表。見出し行は紺の地に白の太字。行のあいだに細い線だけ。highlight_col の列は淡い青の地で濃い青の太字。"""
        nr, nc = len(rows), len(rows[0])
        row_h = row_h or int(Pt(size) * 2.1)
        gf = slide.shapes.add_table(nr, nc, box.x, box.y, box.w, row_h * nr)
        tbl = gf.table
        tblPr = gf._element.graphic.graphicData.tbl.tblPr
        tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        col_w = col_w or [box.w // nc] * nc
        for j, w in enumerate(col_w):
            tbl.columns[j].width = int(w)
        for i in range(nr):
            tbl.rows[i].height = row_h
            for j in range(nc):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = 91440
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                hl = highlight_col == j
                if i == 0:
                    c.fill.solid(); c.fill.fore_color.rgb = BLUE_DARK if hl else NAVY
                elif hl:
                    c.fill.solid(); c.fill.fore_color.rgb = PALE
                else:
                    c.fill.background()
                p = c.text_frame.paragraphs[0]
                p.alignment = (align[j] if align else PP_ALIGN.CENTER)
                if i == 0:
                    self._runs(p, rows[i][j], size=size - 2, bold=True, color=WHITE)
                else:
                    b = hl or i in bold_rows
                    self._runs(p, rows[i][j], size=size, bold=b, color=BLUE_DARK if b else TEXT)
                tcPr = c._tc.get_or_add_tcPr()
                for side in ("a:lnL", "a:lnR", "a:lnT"):
                    ln = etree.SubElement(tcPr, qn(side), w="0"); etree.SubElement(ln, qn("a:noFill"))
                lnb = etree.SubElement(tcPr, qn("a:lnB"), w=str(int(Pt(0.75))))
                sf = etree.SubElement(lnb, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr"), val="D9DEE6")
                for tag in ("a:solidFill", "a:noFill"):
                    fill = tcPr.find(qn(tag))
                    if fill is not None and fill is not tcPr[-1]:
                        tcPr.remove(fill); tcPr.append(fill)
        return gf

    def image(self, slide, box, path=None, label="画像", caption=None):
        """角丸の画像。path が無いときは破線の差し込み枠。caption は下に灰色の小さい文字で。"""
        if path and os.path.exists(path):
            # 縦横比を保ったまま box に収め、中央に置く
            from PIL import Image
            iw, ih = Image.open(path).size
            sc = min(box.w / iw, box.h / ih)
            pw, ph = int(iw * sc), int(ih * sc)
            shp = slide.shapes.add_picture(path, box.x + (box.w - pw) // 2, box.y + (box.h - ph) // 2, pw, ph)
            shp._element.spPr.find(qn("a:prstGeom")).set("prst", "roundRect")
            avl = shp._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
            if avl is not None:
                etree.SubElement(avl, qn("a:gd"), name="adj", fmla="val 4000")
            box = Box(box.x, box.y + (box.h - ph) // 2, box.w, ph)
        else:
            shp = _rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, box, fill=RGBColor(0xF5, 0xF6, 0xF8), line=GRAY,
                        line_w=1.25, radius=0.06)
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for j, t in enumerate(["［画像を差し込む］", label]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = t; _font(r, SZ_SMALL - 2, False, GRAY)
            shp.name = f"画像の差し込み位置：{label}"
        if caption:
            self.text(slide, Box(box.x, box.b + int(0.08 * IN), box.w, int(0.35 * IN)), [caption], size=SZ_SMALL - 2,
                      color=GRAY, align=PP_ALIGN.CENTER)
        return shp

    def references(self, slide, lines):
        """参考文献と注。下端の表示のすぐ上に小さく。"""
        h = int(Pt(SZ_FOOT) * 1.35) * len(lines)
        return self.text(slide, Box(int(0.6 * IN), int(6.92 * IN) - h, int(12.1 * IN), h), lines, size=SZ_FOOT,
                         color=GRAY, line_spacing=1.0)

    def text(self, slide, box, lines, size=SZ_BODY, bold=False, color=None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.2, space_after=0):
        """自由配置のテキスト。lines は文字列か run のリストのリスト。改行位置は呼び出し側で決める。"""
        tb = slide.shapes.add_textbox(box.x, box.y, box.w, box.h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            if space_after:
                p.space_after = Pt(space_after)
            self._runs(p, line, size=size, bold=bold, color=color)
        return tb

    def save(self, path, author="", title=""):
        """保存する。テンプレートの元の資料の情報（作成者、日時、統計、サムネイル）は残さず、
        この資料の情報に置き換える。サムネイルは 1 枚目を描いた画像にする（LibreOffice が無ければ外す）。"""
        import datetime
        n = len(self.prs.slides)
        cp = self.prs.core_properties
        now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None, microsecond=0)
        cp.author, cp.last_modified_by, cp.title = author, author, title
        cp.created, cp.modified, cp.revision = now, now, 1
        index = {id(s): i for i, s in enumerate(self.prs.slides)}
        for s, num in self._pages:
            k = index[id(s)] + 1
            p = num.text_frame.paragraphs[0]
            for r in list(p.runs):
                p._p.remove(r._r)
            self._runs(p, [(str(k), {"bold": True, "color": BLUE}), (f" / {n}", {"bold": False, "color": GRAY})],
                       size=SZ_FOOT + 1)
        self.prs.save(path)
        _clean_package(path, n)


    def _runs(self, p, text, size=None, bold=False, color=None):
        """text は文字列、または (文字列, 書式) のリスト。
        書式: bold, size, color（RGBColor か "blue" か "key"）, sup（上付き）, sub（下付き）
        "key" は結論の最重要語：濃い青の太字に黄色の太い下線。"""
        parts = [(text, {})] if isinstance(text, str) else text
        for t, fmt in parts:
            r = p.add_run()
            r.text = t
            col = fmt.get("color", color)
            key = col == "key"
            if isinstance(col, str):
                col = {"blue": BLUE_DARK, "key": BLUE_DARK}[col]
            _font(r, fmt.get("size", size), True if key else fmt.get("bold", bold), col)
            rPr = r._r.get_or_add_rPr()
            if key:
                rPr.set("u", "heavy")
                uf = etree.Element(qn("a:uFill"))
                sf = etree.SubElement(uf, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr"), val=YELLOW)
                rPr.find(qn("a:latin")).addprevious(uf)
            if fmt.get("sup"):
                rPr.set("baseline", "30000")
            if fmt.get("sub"):
                rPr.set("baseline", "-25000")


def _render_first_slide(path):
    """1 枚目を JPEG に描く。LibreOffice が無い、または失敗したときは None。"""
    import glob
    import shutil
    import subprocess
    import tempfile
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out = tempfile.mkdtemp()
    try:
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{out}/profile",
                        "--convert-to", "jpg", "--outdir", out, path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        jpg = glob.glob(os.path.join(out, "*.jpg"))
        if not jpg:
            return None
        from PIL import Image
        import io
        im = Image.open(jpg[0]).convert("RGB")
        im.thumbnail((256, 256))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
        return buf.getvalue()
    except Exception:
        return None
    finally:
        shutil.rmtree(out, ignore_errors=True)


def _clean_package(path, n_slides):
    """テンプレート由来の文書情報を消す：統計（app.xml）を作り直し、独自の属性（custom.xml）を外し、
    サムネイルを 1 枚目の画像に差し替える（描けなければサムネイルごと外す）。"""
    import re
    import shutil
    import zipfile
    thumb = _render_first_slide(path)
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f'<Application>Microsoft Office PowerPoint</Application><Slides>{n_slides}</Slides>'
           '</Properties>')
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            data = zin.read(name)
            if name == "docProps/custom.xml":
                continue
            if name == "docProps/thumbnail.jpeg":
                if thumb is None:
                    continue
                data = thumb
            elif name == "docProps/app.xml":
                data = app.encode("utf-8")
            elif name == "_rels/.rels":
                t = data.decode("utf-8")
                t = re.sub(r'<Relationship [^>]*Target="docProps/custom.xml"/>', "", t)
                if thumb is None:
                    t = re.sub(r'<Relationship [^>]*Target="docProps/thumbnail.jpeg"/>', "", t)
                data = t.encode("utf-8")
            elif name == "[Content_Types].xml":
                t = data.decode("utf-8")
                t = re.sub(r'<Override [^>]*PartName="/docProps/custom.xml"[^>]*/>', "", t)
                data = t.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
